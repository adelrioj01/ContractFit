from io import BytesIO

import pymupdf
import pytest
from docx import Document

from contractfit.document_parser import DocumentParseError, DocumentParser

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def _two_page_pdf() -> bytes:
    document = pymupdf.open()
    first_page = document.new_page()
    first_page.insert_text((72, 72), "First page has enough searchable text for parsing.")
    second_page = document.new_page()
    second_page.insert_text((72, 72), "scan")
    data = document.tobytes()
    document.close()
    return data


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("Solicitation title")
    document.add_paragraph("Mandatory site visit")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_parses_pdf_by_page_and_marks_short_pages_for_ocr() -> None:
    pages = DocumentParser().parse("solicitation.pdf", "application/pdf", _two_page_pdf())

    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert "First page" in pages[0].text
    assert pages[0].needs_ocr is False
    assert pages[1].page_number == 2
    assert "scan" in pages[1].text
    assert pages[1].needs_ocr is True


def test_parses_docx_as_one_page() -> None:
    pages = DocumentParser().parse("solicitation.docx", DOCX_CONTENT_TYPE, _docx())

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].text == "Solicitation title\nMandatory site visit"
    assert pages[0].needs_ocr is False


def test_parses_utf8_plain_text() -> None:
    pages = DocumentParser().parse(
        "notice.txt", "text/plain; charset=utf-8", "Oportunidad pública".encode()
    )

    assert len(pages) == 1
    assert pages[0].text == "Oportunidad pública"


@pytest.mark.parametrize(
    ("filename", "content_type", "data"),
    [
        ("image.png", "image/png", b"not an image"),
        ("broken.pdf", "application/pdf", b"not a pdf"),
        ("broken.docx", DOCX_CONTENT_TYPE, b"not a docx"),
        ("broken.txt", "text/plain", b"\xff"),
    ],
)
def test_rejects_unsupported_or_malformed_documents(
    filename: str, content_type: str, data: bytes
) -> None:
    with pytest.raises(DocumentParseError):
        DocumentParser().parse(filename, content_type, data)
